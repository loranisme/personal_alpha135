"""Human review and immutable freezing for a small active alpha pool."""

from __future__ import annotations

import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from collections.abc import Mapping
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


DECISIONS = ("INCLUDE", "WATCHLIST", "REJECT")
CONFLICT_COLUMNS = (
    "factor_id_a",
    "factor_id_b",
    "absolute_score_correlation",
    "top_decile_overlap",
    "status",
    "universe_label",
    "sample_start",
    "sample_end",
    "provider",
    "input_hash",
)
_HEX64 = re.compile(r"[0-9a-f]{64}", re.IGNORECASE)


def _validate_policy(policy: Mapping[str, Any]) -> None:
    if int(policy.get("minimum_factors", 0)) != 6 or int(policy.get("maximum_factors", 0)) != 8:
        raise ValueError("INVALID_ACTIVE_POOL_SIZE_POLICY")
    if policy.get("default_provider") != "tiingo_eod":
        raise ValueError("UNSUPPORTED_ACTIVE_POOL_PROVIDER")
    correlation = float(policy.get("absolute_score_correlation_warning", -1))
    overlap = float(policy.get("top_decile_overlap_warning", -1))
    if not 0 < correlation < 1 or not 0 < overlap < 1:
        raise ValueError("INVALID_CONFLICT_POLICY")


def build_review_candidates(
    catalog: pd.DataFrame,
    coverage: pd.DataFrame,
    policy: Mapping[str, Any],
) -> pd.DataFrame:
    """Return every default-eligible row in deterministic structural order."""
    _validate_policy(policy)
    if set(coverage.columns) != {"local_factor_id", "coverage"}:
        raise ValueError("INVALID_COVERAGE_SCHEMA")
    if coverage.local_factor_id.duplicated().any():
        raise ValueError("DUPLICATE_COVERAGE_FACTOR_ID")
    required = {
        "local_factor_id",
        "default_provider_eligible",
        "allowed_providers",
        "mechanism_tag",
        "semantic_status",
        "warmup_sessions",
    }
    if not required.issubset(catalog.columns):
        raise ValueError("INVALID_CATALOG_SCHEMA")
    eligible = catalog.loc[catalog.default_provider_eligible == True].copy()  # noqa: E712
    eligible = eligible.merge(coverage, on="local_factor_id", how="left", validate="one_to_one")
    if eligible.coverage.isna().any():
        raise ValueError("CANDIDATE_COVERAGE_MISSING")
    eligible["dual_provider"] = (
        eligible.allowed_providers.astype(str).str.contains("alpaca_sip", regex=False)
        & eligible.allowed_providers.astype(str).str.contains("tiingo_eod", regex=False)
    )
    eligible["semantic_order"] = eligible.semantic_status.map(
        {"FULL_INTENT": 0, "PARTIAL_INTENT": 1}
    ).fillna(2)
    eligible = eligible.sort_values(
        [
            "mechanism_tag",
            "dual_provider",
            "coverage",
            "semantic_order",
            "warmup_sessions",
            "local_factor_id",
        ],
        ascending=[True, False, False, True, True, True],
        kind="stable",
    )
    eligible["mechanism_rank"] = eligible.groupby(
        "mechanism_tag", sort=False
    ).cumcount() + 1
    eligible["structural_recommendation"] = eligible.mechanism_rank.eq(1)
    return eligible.drop(columns="semantic_order").reset_index(drop=True)


def _top_decile_overlap(left: pd.Series, right: pd.Series) -> float:
    aligned = pd.concat([left.rename("left"), right.rename("right")], axis=1).dropna()
    if aligned.empty:
        return float("nan")
    if isinstance(aligned.index, pd.MultiIndex):
        groups = (frame for _, frame in aligned.groupby(level=0, sort=True))
    else:
        groups = (aligned,)
    values = []
    for frame in groups:
        k = max(1, math.ceil(len(frame) * 0.10))
        left_top = set(frame.nlargest(k, "left").index)
        right_top = set(frame.nlargest(k, "right").index)
        values.append(len(left_top & right_top) / k)
    return float(sum(values) / len(values))


def _conflicts(
    candidates: pd.DataFrame,
    scores: Mapping[str, pd.Series] | None,
    policy: Mapping[str, Any],
) -> pd.DataFrame:
    evidence = dict(policy.get("score_evidence") or {})
    base_evidence = {
        "universe_label": evidence.get("universe_label", "NOT_SUPPLIED"),
        "sample_start": evidence.get("sample_start"),
        "sample_end": evidence.get("sample_end"),
        "provider": evidence.get("provider"),
        "input_hash": evidence.get("input_hash"),
    }
    if not scores:
        return pd.DataFrame(
            [{"status": "NOT_SUPPLIED", **base_evidence}], columns=CONFLICT_COLUMNS
        )
    rows: list[dict[str, Any]] = []
    factor_ids = [item for item in candidates.local_factor_id if item in scores]
    for index, left_id in enumerate(factor_ids):
        for right_id in factor_ids[index + 1 :]:
            aligned = pd.concat(
                [scores[left_id].rename("left"), scores[right_id].rename("right")], axis=1
            ).dropna()
            correlation = (
                abs(float(aligned.left.corr(aligned.right, method="spearman")))
                if len(aligned) >= 2
                else float("nan")
            )
            overlap = _top_decile_overlap(scores[left_id], scores[right_id])
            if (
                math.isfinite(correlation)
                and correlation > float(policy["absolute_score_correlation_warning"])
            ) or (
                math.isfinite(overlap)
                and overlap > float(policy["top_decile_overlap_warning"])
            ):
                rows.append(
                    {
                        "factor_id_a": left_id,
                        "factor_id_b": right_id,
                        "absolute_score_correlation": correlation,
                        "top_decile_overlap": overlap,
                        "status": "CONFLICT",
                        **base_evidence,
                    }
                )
    if not rows:
        rows.append({"status": "NO_CONFLICTS", **base_evidence})
    return pd.DataFrame(rows, columns=CONFLICT_COLUMNS)


def _cell(value: Any) -> Any:
    if value is None or value is pd.NA:
        return None
    try:
        if bool(pd.isna(value)):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    if hasattr(value, "item"):
        return value.item()
    return value


def _write_frame(sheet, frame: pd.DataFrame) -> None:
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")


def build_active_pool_review(
    catalog: pd.DataFrame,
    coverage: pd.DataFrame,
    scores: Mapping[str, pd.Series] | None,
    policy: Mapping[str, Any],
    output_path: Path | str,
) -> dict[str, Any]:
    """Write the protected three-sheet workbook used for human selection."""
    output = Path(output_path)
    if output.exists():
        raise FileExistsError("ACTIVE_POOL_REVIEW_EXISTS")
    candidates = build_review_candidates(catalog, coverage, policy)
    conflicts = _conflicts(candidates, scores, policy)
    decisions = candidates.copy()
    decisions["decision"] = "WATCHLIST"
    decisions["comment"] = ""

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in (
        ("Family Shortlist", candidates),
        ("Correlation Conflicts", conflicts),
        ("Decision Template", decisions),
    ):
        sheet = workbook.create_sheet(name)
        _write_frame(sheet, frame)
        sheet.protection.sheet = True
        if name == "Decision Template":
            headers = {cell.value: cell.column for cell in sheet[1]}
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, headers["decision"]).protection = Protection(locked=False)
                sheet.cell(row, headers["comment"]).protection = Protection(locked=False)
            validation = DataValidation(
                type="list", formula1='"INCLUDE,WATCHLIST,REJECT"', allow_blank=False
            )
            sheet.add_data_validation(validation)
            validation.add(
                f"{sheet.cell(2, headers['decision']).coordinate}:"
                f"{sheet.cell(sheet.max_row, headers['decision']).coordinate}"
            )
            decision_letter = sheet.cell(1, headers["decision"]).column_letter
            sheet.conditional_formatting.add(
                f"{decision_letter}2:{decision_letter}{sheet.max_row}",
                FormulaRule(
                    formula=[f'${decision_letter}2="INCLUDE"'],
                    fill=PatternFill("solid", fgColor="C6EFCE"),
                ),
            )
    workbook.save(output)
    reopened = load_workbook(output)
    if reopened.sheetnames != ["Family Shortlist", "Correlation Conflicts", "Decision Template"]:
        raise ValueError("ACTIVE_POOL_WORKBOOK_SHEETS_MISMATCH")
    decision_sheet = reopened["Decision Template"]
    if decision_sheet.max_row - 1 != len(candidates):
        raise ValueError("ACTIVE_POOL_WORKBOOK_ROW_MISMATCH")
    if len(decision_sheet.data_validations.dataValidation) != 1:
        raise ValueError("ACTIVE_POOL_WORKBOOK_VALIDATION_MISSING")
    return {
        "status": "PASS",
        "path": output,
        "candidate_count": len(candidates),
        "conflict_count": int(conflicts.status.eq("CONFLICT").sum()),
        "qa": {"reopened": True, "sheet_names": reopened.sheetnames},
    }


def _sheet_frame(workbook, name: str) -> pd.DataFrame:
    sheet = workbook[name]
    values = list(sheet.values)
    if not values:
        return pd.DataFrame()
    return pd.DataFrame(values[1:], columns=values[0]).dropna(how="all")


def _canonical_hash(value: Mapping[str, Any]) -> str:
    encoded = json.dumps(
        value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def freeze_active_pool(
    review_path: Path | str,
    library: Mapping[str, Any],
    manifest_sha256: str,
    policy: Mapping[str, Any],
    output_path: Path | str,
) -> dict[str, Any]:
    """Validate human decisions and freeze a non-overwritable Active Pool."""
    _validate_policy(policy)
    if not _HEX64.fullmatch(str(manifest_sha256)):
        raise ValueError("INVALID_V4_MANIFEST_SHA256")
    output = Path(output_path)
    if output.exists():
        raise FileExistsError("ACTIVE_POOL_OUTPUT_EXISTS")
    workbook = load_workbook(review_path, data_only=False)
    if workbook.sheetnames != ["Family Shortlist", "Correlation Conflicts", "Decision Template"]:
        raise ValueError("ACTIVE_POOL_WORKBOOK_SHEETS_MISMATCH")
    candidates = _sheet_frame(workbook, "Family Shortlist")
    conflicts = _sheet_frame(workbook, "Correlation Conflicts")
    decisions = _sheet_frame(workbook, "Decision Template")
    if candidates.direction.tolist() != decisions.direction.tolist():
        raise ValueError("DIRECTION_MUST_REMAIN_ONE")
    structural = [column for column in candidates.columns if column in decisions.columns]
    if candidates.loc[:, structural].reset_index(drop=True).to_dict("records") != decisions.loc[:, structural].reset_index(drop=True).to_dict("records"):
        raise ValueError("READ_ONLY_REVIEW_FIELDS_CHANGED")
    if not decisions.decision.isin(DECISIONS).all():
        raise ValueError("INVALID_HUMAN_DECISION")
    selected = decisions.loc[decisions.decision.eq("INCLUDE")].copy()
    minimum = int(policy["minimum_factors"])
    maximum = int(policy["maximum_factors"])
    if not minimum <= len(selected) <= maximum:
        raise ValueError("ACTIVE_POOL_SIZE")
    if selected.local_factor_id.duplicated().any():
        raise ValueError("DUPLICATE_LOCAL_FACTOR_ID")
    if selected.comment.isna().any() or selected.comment.astype(str).str.strip().eq("").any():
        raise ValueError("INCLUDE_COMMENT_REQUIRED")

    indexed = {factor.get("local_factor_id"): factor for factor in library.get("factors", [])}
    selected_rows: list[dict[str, Any]] = []
    for row in selected.to_dict("records"):
        factor_id = row["local_factor_id"]
        factor = indexed.get(factor_id)
        if factor is None:
            raise ValueError(f"FROZEN_FACTOR_ID_MISSING:{factor_id}")
        mechanism_tags = (factor.get("economic_description") or {}).get("mechanism_tags") or []
        mechanism = mechanism_tags[0] if mechanism_tags else "unclassified"
        if int(row.get("direction")) != 1:
            raise ValueError("DIRECTION_MUST_REMAIN_ONE")
        if row.get("family_id") != factor.get("family_id") or row.get("mechanism_tag") != mechanism:
            raise ValueError("READ_ONLY_REVIEW_FIELDS_CHANGED")
        if (
            factor.get("build_status") != "COMPUTE_VERIFIED"
            or "tiingo_eod" not in factor.get("allowed_providers", [])
            or mechanism == "unclassified"
        ):
            raise ValueError("FACTOR_NOT_DEFAULT_PROVIDER_ELIGIBLE")
        selected_rows.append(
            {
                "local_factor_id": factor_id,
                "family_id": factor.get("family_id"),
                "mechanism_tag": mechanism,
                "direction": 1,
                "comment": str(row["comment"]).strip(),
            }
        )
    families = [row["family_id"] for row in selected_rows]
    if len(set(families)) != len(families):
        raise ValueError("DUPLICATE_FAMILY")
    selected_ids = {row["local_factor_id"] for row in selected_rows}
    for conflict in conflicts.to_dict("records"):
        if (
            conflict.get("status") == "CONFLICT"
            and conflict.get("factor_id_a") in selected_ids
            and conflict.get("factor_id_b") in selected_ids
        ):
            raise ValueError("UNRESOLVED_FACTOR_CONFLICT")

    mechanisms: dict[str, list[str]] = defaultdict(list)
    for row in selected_rows:
        mechanisms[row["mechanism_tag"]].append(row["local_factor_id"])
    warnings = [
        {
            "code": "SAME_MECHANISM_INCLUDED",
            "mechanism_tag": mechanism,
            "factor_ids": sorted(factor_ids),
        }
        for mechanism, factor_ids in sorted(mechanisms.items())
        if len(factor_ids) > 1
    ]
    with localcontext() as context:
        context.prec = 28
        weight = str(Decimal(1) / Decimal(len(selected_rows)))
    for row in selected_rows:
        row["weight"] = weight
    decision_list = [
        {
            "local_factor_id": str(row["local_factor_id"]),
            "decision": str(row["decision"]),
            "comment": "" if row.get("comment") is None else str(row.get("comment")),
        }
        for row in decisions.to_dict("records")
    ]
    identity = {
        "v4_manifest_sha256": manifest_sha256,
        "provider": "tiingo_eod",
        "factor_ids": [row["local_factor_id"] for row in selected_rows],
        "comments": [row["comment"] for row in selected_rows],
    }
    active_pool_id = "active_pool__" + _canonical_hash(identity)[:16]
    payload: dict[str, Any] = {
        "schema_version": 1,
        "status": "FROZEN_WITH_WARNINGS" if warnings else "FROZEN",
        "active_pool_id": active_pool_id,
        "library_version": "reconstruction-v4",
        "v4_manifest_sha256": manifest_sha256,
        "provider": "tiingo_eod",
        "selected_factors": selected_rows,
        "decision_list": decision_list,
        "warnings": warnings,
        "selection_basis": "HUMAN_REVIEW_WITH_DIAGNOSTIC_ONLY_REFERENCE",
        "data_exposure": dict(policy.get("score_evidence") or {}),
        "live_orders_submitted": 0,
    }
    payload["content_sha256"] = _canonical_hash(payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("x", encoding="utf-8") as stream:
        stream.write(json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n")
    reopened = json.loads(output.read_text(encoding="utf-8"))
    declared = reopened.pop("content_sha256")
    if declared != _canonical_hash(reopened):
        raise ValueError("ACTIVE_POOL_REOPEN_HASH_MISMATCH")
    return payload
