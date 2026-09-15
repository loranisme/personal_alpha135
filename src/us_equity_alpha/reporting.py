"""Immutable local two-stage reports; workbook values are verified after reopening."""
from __future__ import annotations
import copy
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
from decimal import Decimal
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def canonical_hash(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,separators=(',',':'),allow_nan=False,default=str).encode()).hexdigest()


def signal_hash(signal):
    return canonical_hash({k:v for k,v in signal.items() if k!='signal_hash'})


def _rows(value):
    if isinstance(value,pd.DataFrame): return value.to_dict('records')
    if isinstance(value,dict): return [{'key':k,'value':v} for k,v in value.items()]
    return [copy.deepcopy(row) if isinstance(row,dict) else {'check':str(row),'status':'UNRESOLVED'} for row in (value or [])]


def _cell(value):
    if value is None or isinstance(value,float) and not math.isfinite(value): return None
    if isinstance(value,(dict,list,tuple)): value=json.dumps(value,sort_keys=True,default=str)
    if isinstance(value,str) and value[:1] in ('=','+','-','@'): return "'"+value
    return value


def _write(bundle,output_root,execution):
    b=copy.deepcopy(bundle)
    signal=b.get('signal_snapshot', b)
    plan=b.get('execution_plan',b)
    sid=signal.get('signal_id',b.get('signal_id'))
    run_id=plan.get('execution_id') if execution else sid
    if not run_id or not re.fullmatch(r'[A-Za-z0-9_.-]+',str(run_id)) or str(run_id) in {'.','..'}: raise ValueError('INVALID_RUN_ID')
    declared=signal.get('signal_hash')
    if (not execution or 'signal_snapshot' in b) and declared and declared!=signal_hash(signal): raise ValueError('SIGNAL_HASH_CONTENT_MISMATCH')
    if execution and plan.get('signal_hash')!=declared: raise ValueError('SIGNAL_HASH_MISMATCH')
    if execution and not declared: raise ValueError('SIGNAL_HASH_REQUIRED')
    if execution and plan.get('release_id',signal.get('release_id'))!=signal.get('release_id'): raise ValueError('RELEASE_MISMATCH')
    ranking=_rows(signal.get('ranking',b.get('ranking')))
    if execution and 'signal_snapshot' in b:
        if 'ranking' in plan and _rows(plan['ranking'])!=ranking: raise ValueError('FROZEN_RANKING_MISMATCH')
        frozen={r.get('security_id',r.get('ticker')):str(r.get('target_weight')) for r in _rows(signal.get('target_weights',signal.get('targets')))}
        for row in _rows(plan.get('target_portfolio',plan.get('targets'))):
            key=row.get('security_id',row.get('ticker'))
            if key in frozen and row.get('target_weight') is not None and Decimal(str(row['target_weight']))!=Decimal(frozen[key]): raise ValueError('FROZEN_WEIGHT_MISMATCH')
    targets=_rows(plan.get('target_portfolio',plan.get('targets',signal.get('target_weights',signal.get('targets'))))) if execution else _rows(signal.get('target_weights',signal.get('targets')))
    orders=_rows(plan.get('orders'))
    engine_plan=copy.deepcopy(plan)
    mode=b.get('mode',plan.get('mode','DEVELOPMENT'))
    evidence=b.get('evidence_status',plan.get('evidence_status','INSUFFICIENT_EVIDENCE'))
    status=plan.get('run_status','BLOCKED') if execution else ('NO_REBALANCE' if signal.get('run_status')=='NO_REBALANCE' else 'SIGNAL_ONLY')
    synthetic=b.get('synthetic_demo',False) or b.get('artifact_kind')=='SYNTHETIC_DEMO'
    reason=[]
    if execution and (synthetic or evidence not in {'VALIDATION_PASS','HOLDOUT_PASS','FORWARD_PENDING'}):
        status='BLOCKED'; reason.append('SYNTHETIC_DEMO_OR_INSUFFICIENT_EVIDENCE')
    if execution and mode=='MANUAL_LIVE' and not b.get('gate_results',{}).get('G6',False):
        status='BLOCKED'; reason.append('G6_NOT_ACTIVATED')
    if status in {'BLOCKED','EXPIRED','SUPERSEDED','NO_REBALANCE'}:
        for row in orders:
            row['approved_trade_shares']=0
            row['order_status']='KEEP' if status=='NO_REBALANCE' else 'BLOCKED'
            row['blocked_reason']=';'.join(reason) or status
            if 'actual_shares' in row: row['executable_target_shares']=row['actual_shares']
        for row in orders+targets:
            for field in ('executable_value','weight_after_rounding','approved_notional','approved_fees','executable_weight'):
                if field in row: row[field]=None
            if row.get('security_id',row.get('ticker'))=='CASH' and 'target_amount' in row: row['target_amount']=None
        for row in targets:
            if 'approved_trade_shares' in row: row['approved_trade_shares']=0
            if 'actual_shares' in row: row['executable_target_shares']=row['actual_shares']
    if not execution:
        forbidden={'approved_trade_shares','trade_shares','target_shares','executable_target_shares','order_status','action'}
        targets=[{k:v for k,v in row.items() if k not in forbidden} for row in targets]
    meta=dict(mode=mode,run_status=status,evidence_status=evidence,artifact_kind='SYNTHETIC_DEMO' if synthetic else 'LOCAL_PRIVATE',signal_hash=declared or signal_hash(signal),units='USD / shares / decimal weights',timezone='UTC unless timestamp carries offset')
    for rows in (ranking,targets,orders):
        for row in rows: row.update(meta)
    checks=[dict(check=k,value=v) for k,v in meta.items()]+_rows(b.get('checks',plan.get('checks')))
    checks += [dict(check='blocking_reason',value=x) for x in reason]
    sheets=[('Ranking',ranking,'stock_ranking.csv'),('Target Portfolio',targets,'target_portfolio.csv' if execution else 'target_weights.csv')]
    if execution: sheets.append(('Rebalance Orders',orders,'rebalance_orders.csv'))
    sheets.append(('Checks',checks,None))
    root=Path(output_root)/('executions' if execution else 'signals'); root.mkdir(parents=True,exist_ok=True)
    final=root/str(run_id)
    if final.exists(): raise FileExistsError(final)
    temp=Path(tempfile.mkdtemp(prefix='.staging-',dir=root))
    try:
        book=Workbook(); book.remove(book.active); expected={}
        for title,rows,csvname in sheets:
            headers=list(dict.fromkeys(k for row in rows for k in row)) or ['status']
            values=[[_cell(row.get(k)) for k in headers] for row in rows]
            sheet=book.create_sheet(title); sheet.append(headers)
            for row in values: sheet.append(row)
            sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='17365D')
            for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=min(45,max(15,len(str(column[0].value))+2))
            expected[title]=[headers]+values
            if csvname:
                with (temp/csvname).open('w',newline='',encoding='utf-8') as f:
                    writer=csv.writer(f); writer.writerow(headers); writer.writerows(values)
        name=f"{'execution' if execution else 'signal'}_{run_id}.xlsx"; book.save(temp/name)
        reopened=load_workbook(temp/name,data_only=False)
        equal=all(list(reopened[title].values)==[tuple(r) for r in rows] for title,rows in expected.items())
        if not equal: raise ValueError('XLSX_REOPEN_MISMATCH')
        for title,_,csvname in sheets:
            if csvname:
                with (temp/csvname).open(newline='',encoding='utf-8') as f: csvrows=list(csv.reader(f))
                normalized=[[str(v) if v is not None else '' for v in r] for r in expected[title]]
                if csvrows!=normalized: raise ValueError('CSV_XLSX_MISMATCH')
        cash=plan.get('summary',{}); cash_ok=None
        if 'cash_after' in cash and all(k in cash for k in ('buy_notional','sell_notional','estimated_fees')):
            initial=Decimal(str(cash['cash_after']))+Decimal(str(cash['buy_notional']))-Decimal(str(cash['sell_notional']))+Decimal(str(cash['estimated_fees']))
            account=b.get('account_snapshot',b.get('account',{}))
            if 'cash' in account: initial=Decimal(str(account['cash']))
            if all(k in cash for k in ('approved_buy_notional','approved_sell_notional','approved_estimated_fees','approved_cash_after')):
                cash_ok=abs(initial-Decimal(str(cash['approved_buy_notional']))+Decimal(str(cash['approved_sell_notional']))-Decimal(str(cash['approved_estimated_fees']))-Decimal(str(cash['approved_cash_after'])))<Decimal('0.0000001')
                if not cash_ok: raise ValueError('CASH_RECONCILIATION_FAILED')
        if all(k in cash for k in ('initial_cash','buy_notional','sell_notional','fees','ending_cash')):
            cash_ok=abs(Decimal(str(cash['initial_cash']))-Decimal(str(cash['buy_notional']))+Decimal(str(cash['sell_notional']))-Decimal(str(cash['fees']))-Decimal(str(cash['ending_cash'])))<Decimal('0.0000001')
            if not cash_ok: raise ValueError('CASH_RECONCILIATION_FAILED')
        qa=dict(csv_excel_equal=equal,cash_tie_out=cash_ok,reopened=True,sheet_names=reopened.sheetnames)
        def dump(filename,obj): (temp/filename).write_text(json.dumps(obj,indent=2,sort_keys=True,allow_nan=False,default=str)+'\n')
        dump('risk_checks.json',dict(**meta,checks=checks,qa=qa))
        if execution:
            dump('engine_intent_audit.json',dict(artifact_kind='ENGINE_CALCULATION_NOT_EXECUTABLE',execution_plan=engine_plan))
            plan.update(run_status=status,orders=orders,target_portfolio=targets,targets=targets,signal_id=sid,signal_hash=declared)
            if status in {'BLOCKED','EXPIRED','SUPERSEDED','NO_REBALANCE'}:
                plan['summary']={'status':'BLOCKED_NO_EXECUTABLE_QUANTITIES','engine_summary_reference':'engine_intent_audit.json'}
            dump('execution_plan.json',plan)
            pd.DataFrame(_rows(b.get('quotes',plan.get('quotes')))).to_parquet(temp/'quotes_snapshot.parquet',index=False)
        else: dump('signal_snapshot.json',signal)
        hashes={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in temp.iterdir() if f.is_file()}
        dump('run_manifest.json',dict(run_id=run_id,signal_id=sid,release_id=signal.get('release_id'),files=hashes,qa=qa,**meta))
        os.rename(temp,final)
        return {('execution_report' if execution else 'signal_report'):final/name,'run_dir':final,'qa':qa,'run_status':status}
    except BaseException:
        shutil.rmtree(temp,ignore_errors=True); raise


def write_signal_report(bundle,output_root): return _write(bundle,output_root,False)
def write_execution_report(bundle,output_root): return _write(bundle,output_root,True)


def write_selection_review(bundle, output_dir):
    """Write and reopen the V5 core selection bundle and optional helper."""
    output=Path(output_dir)
    if output.exists(): raise FileExistsError('OUTPUT_DIRECTORY_EXISTS')
    output.mkdir(parents=True)
    tables=[('data_checks','Data Checks','data_checks.csv'),('alpha_scores','Alpha Scores','alpha_scores.parquet'),('stock_ranking','Stock Ranking','stock_ranking.csv'),('target_portfolio','Target Portfolio','target_portfolio.csv'),('blocked_items','Blocked Items','blocked_items.csv')]
    if bundle.get('execution_helper_status')=='EXECUTION_HELPER_READY': tables.append(('manual_rebalance_draft','Manual Rebalance Draft','manual_rebalance_draft.csv'))
    paths={}
    with pd.ExcelWriter(output/'selection_review.xlsx',engine='openpyxl') as writer:
        for key,sheet,name in tables:
            frame=bundle.get(key,pd.DataFrame())
            path=output/name
            if path.suffix=='.parquet': frame.to_parquet(path,index=False)
            else: frame.to_csv(path,index=False)
            frame.to_excel(writer,sheet_name=sheet,index=False)
            paths[key]=path
    workbook=load_workbook(output/'selection_review.xlsx')
    expected=[x[1] for x in tables]
    if workbook.sheetnames!=expected: raise ValueError('SELECTION_WORKBOOK_SHEETS_MISMATCH')
    for key,sheet,name in tables:
        path=output/name
        frame=pd.read_parquet(path) if path.suffix=='.parquet' else pd.read_csv(path)
        if workbook[sheet].max_row-1 != len(frame): raise ValueError('SELECTION_ARTIFACT_ROW_MISMATCH')
    paths['workbook']=output/'selection_review.xlsx'
    files={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in output.iterdir() if p.is_file()}
    manifest={'schema_version':1,'status':bundle.get('status'),'execution_helper_status':bundle.get('execution_helper_status','NOT_REQUESTED'),'files':files,'live_orders_submitted':0}
    (output/'manifest.json').write_text(json.dumps(manifest,indent=2,sort_keys=True)+'\n')
    paths['manifest']=output/'manifest.json'
    return paths
