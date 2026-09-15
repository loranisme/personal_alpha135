import json
import pandas as pd
from openpyxl import load_workbook
from us_equity_alpha.reporting import write_selection_review

def test_selection_writer_conditionally_includes_helper(tmp_path):
    ranking=pd.DataFrame([{"security_id":"A","composite_score":.9}])
    targets=pd.DataFrame([{"security_id":"A","target_weight":.95},{"security_id":"CASH","target_weight":.05}])
    base={"status":"SELECTION_READY","stock_ranking":ranking,"target_portfolio":targets,"alpha_scores":pd.DataFrame([{"security_id":"A","F1":1.0}]),"data_checks":pd.DataFrame([{"check":"coverage","status":"PASS"}]),"blocked_items":pd.DataFrame(columns=["code"]),"execution_helper_status":"NOT_REQUESTED"}
    paths=write_selection_review(base,tmp_path/"core")
    assert not (tmp_path/"core/manual_rebalance_draft.csv").exists()
    assert "Manual Rebalance Draft" not in load_workbook(paths["workbook"]).sheetnames
    manifest=json.loads(paths["manifest"].read_text())
    assert manifest["execution_helper_status"]=="NOT_REQUESTED" and manifest["live_orders_submitted"]==0
    with_helper=dict(base,execution_helper_status="EXECUTION_HELPER_READY",manual_rebalance_draft=pd.DataFrame([{"security_id":"A","suggested_trade_shares":2,"review_status":"REVIEW_REQUIRED"}]))
    paths=write_selection_review(with_helper,tmp_path/"helper")
    assert (tmp_path/"helper/manual_rebalance_draft.csv").is_file()
    assert "Manual Rebalance Draft" in load_workbook(paths["workbook"]).sheetnames
