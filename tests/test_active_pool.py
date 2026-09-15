import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from us_equity_alpha.active_pool import (
    build_active_pool_review,
    build_review_candidates,
    freeze_active_pool,
)
from us_equity_alpha.alpha_catalog import build_alpha_catalog, build_catalog_review_view
from us_equity_alpha.cli import main


MECHANISMS = (
    "momentum",
    "momentum",
    "mean_reversion",
    "realized_risk",
    "trading_activity",
    "intraday_return",
    "range_position",
    "overnight_return",
)


def _factor(index: int, *, family: str | None = None) -> dict:
    factor_id = f"F{index}"
    return {
        "local_factor_id": factor_id,
        "expression": "rank(ts_delta(close, 5))",
        "expression_hash": f"{index + 1:064x}",
        "family_id": family or f"family-{index}",
        "economic_description": {"mechanism_tags": [MECHANISMS[index]]},
        "semantic_status": "FULL_INTENT" if index % 2 == 0 else "PARTIAL_INTENT",
        "field_bindings": {"close": "close"},
        "allowed_providers": ["alpaca_sip", "tiingo_eod"],
        "direction_basis": "Signed source component; unvalidated locally.",
        "build_status": "COMPUTE_VERIFIED",
        "source_alpha_ids": [f"brain-{index}"],
        "settings": {"delay": 1, "decay": 0, "neutralization": "NONE"},
    }


def _library(*, duplicate_family: bool = False) -> dict:
    factors = [_factor(index) for index in range(8)]
    if duplicate_family:
        factors[1]["family_id"] = factors[0]["family_id"]
    return {"schema_version": 2, "factors": factors}


def _catalog_policy() -> dict:
    return {
        "schema_version": 1,
        "library_version": "reconstruction-v4",
        "default_provider": "tiingo_eod",
        "diagnostic_value_columns": [
            "diagnostic_rank_ic_mean_5d",
            "diagnostic_rank_icir_5d",
            "diagnostic_top_bottom_net_spread_5d",
            "diagnostic_turnover_mean",
        ],
    }


def _pool_policy() -> dict:
    return {
        "schema_version": 1,
        "minimum_factors": 6,
        "maximum_factors": 8,
        "default_provider": "tiingo_eod",
        "absolute_score_correlation_warning": 0.80,
        "top_decile_overlap_warning": 0.70,
        "score_evidence": {
            "universe_label": "ENGINEERING_50",
            "sample_start": "2020-01-02",
            "sample_end": "2025-12-31",
            "provider": "tiingo_eod",
            "input_hash": "a" * 64,
        },
    }


def _catalog(library: dict | None = None) -> pd.DataFrame:
    base = build_alpha_catalog(
        library or _library(), _catalog_policy(), library_version="reconstruction-v4"
    )
    view = build_catalog_review_view(base, None, _catalog_policy())
    view["diagnostic_rank_ic_mean_5d"] = np.arange(len(view), dtype=float) / 100
    return view


def _coverage() -> pd.DataFrame:
    return pd.DataFrame(
        {"local_factor_id": [f"F{i}" for i in range(8)], "coverage": [0.99] * 8}
    )


def _scores(*, conflict: bool = False) -> dict[str, pd.Series]:
    index = pd.MultiIndex.from_product(
        [pd.date_range("2025-01-02", periods=3, tz="UTC"), [f"S{i}" for i in range(20)]],
        names=["date", "security_id"],
    )
    output = {}
    base = np.arange(len(index), dtype=float)
    for factor_index in range(8):
        values = np.sin(base * (factor_index + 1) + factor_index)
        output[f"F{factor_index}"] = pd.Series(values, index=index)
    if conflict:
        output["F1"] = output["F0"].copy()
    return output


def _include_first_six(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Decision Template"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        decision = "INCLUDE" if row <= 7 else "WATCHLIST"
        sheet.cell(row, headers["decision"], decision)
        sheet.cell(row, headers["comment"], f"manual review {row}")
    workbook.save(path)


def test_candidate_list_keeps_all_eligible_rows_and_ignores_diagnostics():
    first = build_review_candidates(_catalog(), _coverage(), _pool_policy())
    changed_catalog = _catalog()
    changed_catalog["diagnostic_rank_ic_mean_5d"] *= -100
    second = build_review_candidates(changed_catalog, _coverage(), _pool_policy())
    assert len(first) == 8
    assert first.local_factor_id.tolist() == second.local_factor_id.tolist()
    assert first.groupby("mechanism_tag").mechanism_rank.min().eq(1).all()


def test_review_workbook_limits_edits_to_decision_and_comment(tmp_path):
    path = tmp_path / "review.xlsx"
    result = build_active_pool_review(_catalog(), _coverage(), None, _pool_policy(), path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Family Shortlist", "Correlation Conflicts", "Decision Template"]
    sheet = workbook["Decision Template"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert result["candidate_count"] == sheet.max_row - 1 == 8
    assert sheet.protection.sheet
    assert not sheet.cell(2, headers["decision"]).protection.locked
    assert not sheet.cell(2, headers["comment"]).protection.locked
    assert sheet.cell(2, headers["local_factor_id"]).protection.locked
    assert len(sheet.data_validations.dataValidation) == 1


def test_same_mechanism_freezes_with_warning_and_equal_weights(tmp_path):
    review = tmp_path / "review.xlsx"
    build_active_pool_review(_catalog(), _coverage(), _scores(), _pool_policy(), review)
    _include_first_six(review)
    output = tmp_path / "active_pool.json"
    result = freeze_active_pool(review, _library(), "b" * 64, _pool_policy(), output)
    assert result["status"] == "FROZEN_WITH_WARNINGS"
    warning = next(row for row in result["warnings"] if row["code"] == "SAME_MECHANISM_INCLUDED")
    assert set(warning["factor_ids"]) == {"F0", "F1"}
    assert len(result["selected_factors"]) == 6
    assert sum(float(row["weight"]) for row in result["selected_factors"]) == pytest.approx(1.0)
    assert json.loads(output.read_text(encoding="utf-8"))["content_sha256"] == result["content_sha256"]


def test_duplicate_family_remains_a_hard_blocker(tmp_path):
    library = _library(duplicate_family=True)
    review = tmp_path / "review.xlsx"
    build_active_pool_review(_catalog(library), _coverage(), _scores(), _pool_policy(), review)
    _include_first_six(review)
    with pytest.raises(ValueError, match="DUPLICATE_FAMILY"):
        freeze_active_pool(review, library, "b" * 64, _pool_policy(), tmp_path / "pool.json")


def test_unresolved_score_conflict_remains_a_hard_blocker(tmp_path):
    review = tmp_path / "review.xlsx"
    build_active_pool_review(_catalog(), _coverage(), _scores(conflict=True), _pool_policy(), review)
    _include_first_six(review)
    with pytest.raises(ValueError, match="UNRESOLVED_FACTOR_CONFLICT"):
        freeze_active_pool(review, _library(), "b" * 64, _pool_policy(), tmp_path / "pool.json")


def test_changed_direction_is_rejected(tmp_path):
    review = tmp_path / "review.xlsx"
    build_active_pool_review(_catalog(), _coverage(), None, _pool_policy(), review)
    _include_first_six(review)
    workbook = load_workbook(review)
    sheet = workbook["Decision Template"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["direction"], -1)
    workbook.save(review)
    with pytest.raises(ValueError, match="DIRECTION_MUST_REMAIN_ONE"):
        freeze_active_pool(review, _library(), "b" * 64, _pool_policy(), tmp_path / "pool.json")


def test_cli_builds_review_then_freezes_human_decisions(tmp_path, capsys):
    catalog_path = tmp_path / "catalog.csv"
    coverage_path = tmp_path / "coverage.csv"
    policy_path = tmp_path / "policy.json"
    library_path = tmp_path / "library.json"
    review_path = tmp_path / "review.xlsx"
    pool_path = tmp_path / "pool.json"
    _catalog().to_csv(catalog_path, index=False)
    _coverage().to_csv(coverage_path, index=False)
    policy_path.write_text(json.dumps(_pool_policy()), encoding="utf-8")
    library_path.write_text(json.dumps(_library()), encoding="utf-8")
    assert main([
        "build-active-pool-review",
        "--catalog", str(catalog_path),
        "--coverage", str(coverage_path),
        "--policy", str(policy_path),
        "--output", str(review_path),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "PASS"
    _include_first_six(review_path)
    assert main([
        "freeze-active-pool",
        "--review", str(review_path),
        "--library", str(library_path),
        "--manifest-sha256", "b" * 64,
        "--policy", str(policy_path),
        "--output", str(pool_path),
    ]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "FROZEN_WITH_WARNINGS"
