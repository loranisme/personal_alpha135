import json
import pytest
from openpyxl import load_workbook
from us_equity_alpha.reporting import write_signal_report, write_execution_report, signal_hash


def bundle():
    signal = dict(signal_id='20260909-001', release_id='demo', ranking=[dict(ticker='AAA', score=1.5, rank=1)], target_weights=[dict(ticker='AAA', target_weight=.5)], cutoff='2026-09-09T20:00:00Z')
    signal['signal_hash'] = signal_hash(signal)
    return dict(signal_snapshot=signal, execution_plan=dict(execution_id='20260910-001', signal_hash=signal['signal_hash'], orders=[dict(ticker='AAA', approved_trade_shares=4, order_status='READY_FOR_REVIEW')], target_portfolio=[dict(ticker='AAA', target_weight=.5)], run_status='READY_FOR_REVIEW'), mode='PAPER', evidence_status='DIAGNOSTIC_ONLY', synthetic_demo=True)


def test_reports_are_atomic_immutable_and_reopen(tmp_path):
    b = bundle()
    s = write_signal_report(b, tmp_path)
    e = write_execution_report(b, tmp_path)
    assert load_workbook(s['signal_report']).sheetnames == ['Ranking','Target Portfolio','Checks']
    assert load_workbook(e['execution_report']).sheetnames == ['Ranking','Target Portfolio','Rebalance Orders','Checks']
    assert e['qa']['csv_excel_equal']
    plan = json.loads((e['run_dir'] / 'execution_plan.json').read_text())
    assert plan['run_status'] == 'BLOCKED'
    assert plan['orders'][0]['approved_trade_shares'] == 0
    with pytest.raises(FileExistsError):
        write_execution_report(b, tmp_path)


def test_mismatch_fails_before_publish(tmp_path):
    b = bundle()
    b['execution_plan']['signal_hash']='wrong'
    with pytest.raises(ValueError, match='SIGNAL_HASH'):
        write_execution_report(b,tmp_path)


def test_missing_value_stays_blank(tmp_path):
    b=bundle(); b['signal_snapshot']['ranking'][0]['score']=None
    b['signal_snapshot']['signal_hash']=signal_hash(b['signal_snapshot'])
    result=write_signal_report(b,tmp_path)
    book=load_workbook(result['signal_report'])
    headers=[c.value for c in book['Ranking'][1]]
    assert book['Ranking'].cell(2,headers.index('score')+1).value is None


def test_frozen_ranking_and_weight_cannot_change(tmp_path):
    b=bundle(); b['execution_plan']['ranking']=[{'ticker':'OTHER'}]
    with pytest.raises(ValueError,match='RANKING'): write_execution_report(b,tmp_path)
    b=bundle(); b['execution_plan']['target_portfolio'][0]['target_weight']=.6
    with pytest.raises(ValueError,match='WEIGHT'): write_execution_report(b,tmp_path)


def test_csv_injection_cash_and_status(tmp_path):
    b=bundle(); b['signal_snapshot']['ranking'][0]['ticker']='=HYPERLINK("bad")'
    b['signal_snapshot']['signal_hash']=signal_hash(b['signal_snapshot'])
    b['execution_plan']['signal_hash']=b['signal_snapshot']['signal_hash']
    b['execution_plan']['summary']=dict(initial_cash='100',buy_notional='20',sell_notional='0',fees='1',ending_cash='79')
    out=write_execution_report(b,tmp_path)
    assert out['qa']['cash_tie_out'] is True
    assert "'=HYPERLINK" in (out['run_dir']/'stock_ranking.csv').read_text()


@pytest.mark.parametrize('status',['EXPIRED','SUPERSEDED','NO_REBALANCE','BLOCKED'])
def test_invalid_batch_never_has_approved_shares(tmp_path,status):
    b=bundle(); b.update(synthetic_demo=False,evidence_status='HOLDOUT_PASS'); b['execution_plan']['run_status']=status
    out=write_execution_report(b,tmp_path)
    assert json.loads((out['run_dir']/'execution_plan.json').read_text())['orders'][0]['approved_trade_shares']==0


def test_saved_signal_hash_stays_valid(tmp_path):
    b=bundle(); out=write_signal_report(b,tmp_path)
    saved=json.loads((out['run_dir']/'signal_snapshot.json').read_text())
    assert saved['signal_hash']==signal_hash(saved)


def test_actual_engine_fixture(tmp_path):
    from pathlib import Path
    b=json.loads(Path('tests/fixtures/daily_bundle.json').read_text())
    b['synthetic_demo']=True
    out=write_execution_report(b,tmp_path)
    assert out['run_status']=='BLOCKED'
    assert out['qa']['cash_tie_out'] is True


@pytest.mark.parametrize('status',['BLOCKED','EXPIRED','SUPERSEDED','NO_REBALANCE'])
def test_invalid_batch_clears_hypothetical_execution_values(tmp_path,status):
    b=bundle(); b.update(synthetic_demo=False,evidence_status='HOLDOUT_PASS')
    b['execution_plan']['run_status']=status
    b['execution_plan']['target_portfolio'] += [dict(security_id='CASH',target_weight='.5',target_amount='123',executable_value='123',weight_after_rounding='.123')]
    b['execution_plan']['target_portfolio'][0]['weight_after_rounding']='.456'
    b['execution_plan']['summary']={'approved_cash_after':'123'}
    out=write_execution_report(b,tmp_path)
    plan=json.loads((out['run_dir']/'execution_plan.json').read_text())
    cash=next(r for r in plan['target_portfolio'] if r.get('security_id')=='CASH')
    assert cash['target_amount'] is None
    assert cash['executable_value'] is None
    assert cash['weight_after_rounding'] is None
    assert cash['target_weight']=='.5'
    assert plan['target_portfolio'][0]['weight_after_rounding'] is None
    assert 'approved_cash_after' not in plan['summary']
    audit=json.loads((out['run_dir']/'engine_intent_audit.json').read_text())
    assert audit['execution_plan']['target_portfolio'][-1]['executable_value']=='123'
